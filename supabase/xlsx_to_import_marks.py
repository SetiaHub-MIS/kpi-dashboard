"""Turn Checklist Mingguan workbooks into a ready-to-run import_marks.sql.

    python supabase/xlsx_to_import_marks.py <workbook.xlsx | folder> [...] [--month 2026-09 | --until 2026-09] [-o out.sql]

Reads every sheet of every workbook given (a folder means every .xlsx in it)
and finds each month block by its WEEK 1..WEEK 4 headings. Works on both
layouts in use: one sheet per person ("ID: BC0010" at the top, or the sheet
named "KP0093_NAME"), and the year workbooks with twelve month blocks stacked
down one sheet. Several people stacked down one sheet also work, as long as
each block has its own "ID:" cell above or beside its WEEK headings.

What it takes from each week column:
  * the SV/AS (or AM) column      -> the scores, one per perkara, in sheet order
  * the name on the NAMA row      -> scored_by, matched to a person at the outlet
  * text on the CATATAN rows      -> the mark's note
A week with no score at all is skipped — not marked yet. The MANAGER column is
not imported; what it holds is copied into the SQL as comments so nothing on
the sheet is lost silently.

Hand-kept sheets drift, so the scores are matched to the form by perkara, not
blindly by row. A row whose name was erased keeps its scores; a score typed on
the CATATAN row under a blank perkara is that perkara's; a block with a stray
or missing row is lined up with the standard form by name. A score that
belongs to no perkara (on a heading, under an already-scored line) is reported
and left out — the workbook's JUMLAH counts those, ours does not. Every such
decision is printed as a note.

The month comes from a "MONTH: SEPTEMBER 2026" label, otherwise from the dates
on the TARIKH row (the month most of them fall in). --month keeps only blocks
for that month — use it on a year workbook to import one month; --until keeps
everything up to and including a month, leaving out one still being marked.

The rows are written into a copy of import_marks.sql, which does every check
(right person, right form length, scorer, confirmed weeks) before saving
anything. Nothing here touches the database: paste the output into the
Supabase SQL editor and run it.
"""

import argparse
import datetime as dt
import difflib
import glob
import re
import sys
import warnings as pywarnings
from collections import Counter
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit('openpyxl is missing; install it with:  python -m pip install openpyxl')

TEMPLATE = Path(__file__).with_name('import_marks.sql')

WEEK_RE = re.compile(r'^WEEK\s*([1-4])$', re.I)
MONTH_RE = re.compile(r'MONTH\s*:?\s*([A-Z]+)\s+(\d{4})', re.I)
# The colon is required: without it a scorer called "Idawati" on the NAMA row
# read as "ID AWATI" and re-filed every later block of the sheet under her.
ID_RE = re.compile(r'^ID\s*:\s*([A-Z0-9]*)$', re.I)
PAYROLL_RE = re.compile(r'^[A-Z]+[0-9][A-Z0-9]*$', re.I)
SHEET_ID_RE = re.compile(r'^([A-Z]+[0-9][A-Z0-9]*)(?=[_\s-]|$)', re.I)
LETTERED_RE = re.compile(r'^[A-Z]\s*\)')
NA = {'N/A', 'NA', '-', '—'}

MONTHS = {name: n for n, names in enumerate([
    (), ('JANUARY', 'JANUARI', 'JAN'), ('FEBRUARY', 'FEBRUARI', 'FEB'), ('MARCH', 'MAC', 'MAR'),
    ('APRIL', 'APR'), ('MAY', 'MEI'), ('JUNE', 'JUN'), ('JULY', 'JULAI', 'JUL'),
    ('AUGUST', 'OGOS', 'AUG', 'OGO'), ('SEPTEMBER', 'SEPT', 'SEP'), ('OCTOBER', 'OKTOBER', 'OCT', 'OKT'),
    ('NOVEMBER', 'NOV'), ('DECEMBER', 'DISEMBER', 'DEC', 'DIS'),
]) for name in names}


def text(v):
    """A cell as trimmed text, or '' for blank."""
    if v is None:
        return ''
    return ' '.join(str(v).split())


def as_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    m = re.match(r'^(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})$', text(v))
    if m:  # typed as text: day first, the way it is written here
        try:
            return dt.date(int(m[3]), int(m[2]), int(m[1]))
        except ValueError:
            return None
    return None


# Headings and placeholders that sit where a scorer's name goes.
NOT_A_NAME = {'SV/AS', 'SV', 'AS', 'AM', 'MANAGER', 'NAMA', 'N/A', 'NA', '-', '.', 'TARIKH', 'TARIKH:', 'PERKARA'}


def name_like(v):
    """Could this cell say who scored? Text that is no number, no date and no
    heading. Dates typed as text come in every shape ('20.3.26', '21-28/2/26'),
    so anything of digits and separators alone is a date; '#VALUE!' is Excel's.
    A payroll number has letters, so it still counts."""
    t = text(v)
    return (bool(t) and not isinstance(v, (int, float, dt.date)) and as_date(v) is None
            and t.upper() not in NOT_A_NAME and not WEEK_RE.match(t)
            and not t.startswith('#') and not re.fullmatch(r'[\d\s/.:\-]+', t))


def sql_text(s):
    return 'NULL' if s is None else "'" + s.replace("'", "''") + "'"


def comment(s):
    return ' '.join(s.split())  # a comment must stay on its line


class Problems(list):
    def add(self, where, why):
        self.append(f'{where}: {why}')


def find_people_and_blocks(ws):
    """Yield (person_id, id_cell, header_row, {week_no: column}, month_label) per block."""
    person, id_cell = None, None
    m = SHEET_ID_RE.match(ws.title.strip())
    if m and not re.fullmatch(r'(SHEET|HELAIAN)\d*', m[1], re.I):  # Excel's own "Sheet1" is not a person
        person, id_cell = m[1].upper(), f"sheet name '{ws.title}'"
    for row in ws.iter_rows():
        weeks, month_label = {}, None
        for i, c in enumerate(row):
            t = text(c.value)
            if not t:
                continue
            idm = ID_RE.match(t)
            if idm:
                found = idm[1] or next((text(x.value) for x in row[i + 1:] if text(x.value)), '')
                if PAYROLL_RE.match(found):
                    person, id_cell = found.upper(), f'{ws.title}!{c.coordinate}'
                continue
            wm = WEEK_RE.match(t)
            if wm:
                weeks[int(wm[1])] = c.column
                continue
            mm = MONTH_RE.search(t)
            if mm and mm[1].upper() in MONTHS:
                month_label = (int(mm[2]), MONTHS[mm[1].upper()])
        if weeks:
            yield person, id_cell, row[0].row, weeks, month_label


def parse_block(ws, header_row, end_row, weeks):
    """The rows of one month block: dates, names, perkara lines, notes."""
    where = lambda col, r: f'{text(ws.title)}!{get_column_letter(col)}{r}'
    cell = lambda r, col: ws.cell(row=r, column=col).value

    # The label column is where PERKARA is written; B on every sheet seen so far.
    label_col = next((col for r in range(header_row, min(end_row, header_row + 6) + 1)
                      for col in range(1, min(weeks.values()))
                      if text(cell(r, col)).upper() == 'PERKARA'), 2)

    dates, manager, nama, elsewhere = {}, {}, {}, {}
    first_line = None
    rows = []  # (row, label, numbered)
    for r in range(header_row + 1, end_row + 1):
        label = text(cell(r, label_col))
        row_head = [text(cell(r, col)).upper() for col in range(1, min(weeks.values()))]
        if label.upper().startswith('JUMLAH'):
            break
        if 'NAMA' in row_head:
            for w, col in weeks.items():
                nama[w] = (r, cell(r, col))
                manager[w] = text(cell(r, col + 1)) or None
                if as_date(cell(r, col)):
                    dates.setdefault(w, as_date(cell(r, col)))
            continue
        if first_line is None:
            if cell(r, 1) == 1 or text(cell(r, 1)) == '1':
                first_line = r
            else:
                for w, col in weeks.items():
                    d = as_date(cell(r, col))
                    if d:
                        dates[w] = d
                    elif name_like(cell(r, col)):
                        elsewhere.setdefault(w, []).append((r, text(cell(r, col))))
                continue
        # A row whose name was erased still holds its perkara's scores; skipping
        # it would drop them and shift nothing else, so it counts as a line.
        scored = any(isinstance(cell(r, col), (int, float)) or re.fullmatch(r'\d+', text(cell(r, col)))
                     for col in weeks.values())
        if label or scored:
            numbered = isinstance(cell(r, 1), (int, float)) or text(cell(r, 1)).isdigit()
            rows.append((r, label, numbered))

    if first_line is None:
        return None

    # A numbered perkara with lettered items under it is a heading, not a line.
    is_item = lambda x: not x[1] or LETTERED_RE.match(x[1])
    lines, headings, notes_rows, category = [], [], [], None
    content = [x for x in rows if not x[1].upper().startswith('CATATAN')]
    for r, label, numbered in rows:
        if label.upper().startswith('CATATAN'):
            # The line directly above, if any: a score typed one row too low lands here.
            above = lines[-1][0] if lines and lines[-1][0] == r - 1 else None
            notes_rows.append((r, category, above))
            continue
        if numbered:
            category = re.sub(r'^\d+\s*', '', label)
            nxt = next((x for x in content if x[0] > r), None)
            if nxt and is_item(nxt):
                headings.append((r, label))
                continue
        lines.append((r, label))

    # Who scored: the NAMA row, unless it holds a date or a heading — rows typed
    # in the wrong order — in which case the one name on the block's other
    # header rows, if there is exactly one.
    scorer, scorer_notes = {}, {}
    for w, col in weeks.items():
        r, v = nama.get(w, (None, None))
        found = elsewhere.get(w, [])
        if name_like(v):
            scorer[w] = text(v)
        elif len(found) == 1:
            scorer[w] = found[0][1]
            scorer_notes[w] = (f"{where(col, found[0][0])}: scorer {found[0][1]!r} taken from this row - "
                               f"the NAMA row holds {text(v)!r}" if text(v) else
                               f"{where(col, found[0][0])}: scorer {found[0][1]!r} taken from this row - NAMA is blank")
        else:
            scorer[w] = None
            if text(v):
                scorer_notes[w] = f"{where(col, r)}: NAMA holds {text(v)!r}, not a name - no scorer recorded"

    return dict(dates=dates, scorer=scorer, scorer_notes=scorer_notes, manager=manager, lines=lines,
                headings=headings, notes_rows=notes_rows, where=where, cell=cell, label_col=label_col)


def week_scores(block, col, problems):
    scores = []
    for r, label in block['lines']:
        v = block['cell'](r, col)
        t = text(v).upper()
        if t == '' or t in NA:
            scores.append(None)
        elif isinstance(v, (int, float)) and float(v).is_integer():
            scores.append(int(v))
        elif re.fullmatch(r'\d+', t):
            scores.append(int(t))
        else:
            problems.add(block['where'](col, r), f'{label}: {text(v)!r} is not a score')
            scores.append(None)
    return scores


def bare_number(v):
    t = text(v)
    if isinstance(v, (int, float)) and float(v).is_integer():
        return int(v)
    return int(t) if re.fullmatch(r'\d+', t) else None


def catatan_scores(block, col, scores, warn):
    """A bare number on a CATATAN row is a score typed one row too low. When
    the perkara directly above is blank that week, it is that perkara's score —
    the workbook's own SUM counts it, so taking it keeps our total equal to the
    sheet's JUMLAH. Otherwise it belongs to nothing and is reported."""
    line_at = {r: i for i, (r, _) in enumerate(block['lines'])}
    for r, _, above in block['notes_rows']:
        n = bare_number(block['cell'](r, col))
        if n is None:
            continue
        here = block['where'](col, r)
        if above is not None and scores[line_at[above]] is None:
            scores[line_at[above]] = n
            warn(f"{here}: {n} typed on the CATATAN row under a blank perkara - "
                 f"taken as {block['lines'][line_at[above]][1]!r}'s score")
        else:
            warn(f"{here}: {n} typed on a CATATAN row - not counted as a score "
                 f"(the workbook's JUMLAH does count it)")


def week_note(block, col, with_category=True):
    """The CATATAN text in one column; bare numbers are scores, handled above."""
    parts = []
    for r, category, _ in block['notes_rows']:
        v = block['cell'](r, col)
        t = text(v)
        if not t or bare_number(v) is not None:
            continue
        parts.append(f'{category}: {t}' if category and with_category else t)
    return '; '.join(parts) or None


def read_workbook(path, only_month, until, problems, warnings, marks, layouts):
    wb = openpyxl.load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        blocks = list(find_people_and_blocks(ws))
        for i, (person, id_cell, header_row, weeks, month_label) in enumerate(blocks):
            end_row = blocks[i + 1][2] - 1 if i + 1 < len(blocks) else ws.max_row
            block = parse_block(ws, header_row, end_row, weeks)
            if block is None:
                continue
            here = f'{path.name} / {text(ws.title)} row {header_row}'

            if month_label:
                period = month_label
            else:
                counted = Counter((d.year, d.month) for d in block['dates'].values())
                period = counted.most_common(1)[0][0] if counted else None
            if (only_month and period != only_month) or (until and period and period > until):
                continue

            for r, label in block['lines']:
                if not label:
                    warnings.append(f"{block['where'](block['label_col'], r)}: row has scores but no perkara "
                                    f"name - kept, and lined up with the form below")
            layout = dict(seq=normalise(block['lines']), labels=[l for _, l in block['lines']],
                          here=here, marks=[])
            layouts.append(layout)

            for w, col in sorted(weeks.items()):
                scores = week_scores(block, col, problems)
                catatan_scores(block, col, scores, warnings.append)
                # The import refuses these (scale_max is 5 at every outlet today);
                # naming the cell here saves hunting for it from the SQL error.
                for (r, label), s in zip(block['lines'], scores):
                    if s is not None and s > 5:
                        warnings.append(f"{block['where'](col, r)}: {label} scored {s} - outside 0 to 5, "
                                        f"the import will refuse it; correct the cell")
                manager_scores = [r for r, _ in block['lines']
                                  if text(block['cell'](r, col + 1)) and text(block['cell'](r, col + 1)).upper() not in NA]
                for r, label in block['headings']:
                    if text(block['cell'](r, col)):
                        warnings.append(f"{block['where'](col, r)}: score on the heading {label!r}, which has "
                                        f"A) B) items under it - not counted (the workbook's JUMLAH does count it)")
                if all(s is None for s in scores):
                    if manager_scores:
                        warnings.append(f'{here} week {w}: MANAGER column has scores but SV/AS is blank, not imported')
                    continue
                if person is None:
                    problems.add(here, 'no payroll number; put "ID: KP0000" above the WEEK headings or name the sheet KP0000_NAME')
                    break
                if period is None:
                    problems.add(here, f'week {w}: cannot tell the month; fill in the TARIKH dates or add "MONTH: SEPTEMBER 2026"')
                    continue
                extra = []
                mgr_name = block['manager'].get(w)
                mgr_note = week_note(block, col + 1, with_category=False)
                if mgr_name or mgr_note or manager_scores:
                    extra.append('MANAGER ' + ', '.join(filter(None, [
                        mgr_name,
                        f'scored {len(manager_scores)} perkara' if manager_scores else None,
                        f'"{mgr_note}"' if mgr_note else None,
                    ])))
                if w in block['scorer_notes']:
                    warnings.append(block['scorer_notes'][w])
                mark = dict(
                    person=person, year=period[0], month=period[1], week=w,
                    scored_by=block['scorer'].get(w), note=week_note(block, col),
                    scores=scores, source=f'{path.name} / {text(ws.title)}', workbook=path.name,
                    id_cell=id_cell, extra=extra,
                )
                marks.append(mark)
                layout['marks'].append(mark)


def normalise(lines):
    return tuple(re.sub(r'[^A-Z]', '', l.upper()) for _, l in lines)


def line_up(layouts, problems, warnings):
    """Line every block up with the standard form, by perkara name.

    The standard forms are the layouts most blocks share. A block with the
    same number of rows keeps its scores in place — a mistyped name ("4",
    "AIR COOLER / CURTAIN") does not move them. A block with a row too many or
    too few is matched name by name: a stray row with no score is dropped, a
    perkara the sheet lacks counts as not marked. Anything else is left for a
    person to fix rather than guessed.
    Returns the standard layouts, for the header of the SQL."""
    counted = Counter(l['seq'] for l in layouts)
    standard = [seq for seq, n in counted.most_common() if n >= 0.1 * len(layouts)]
    labels_of = {}
    for l in layouts:
        labels_of.setdefault(l['seq'], l)
    # A form line ending in ':' ("D) LAIN-LAIN:") is there to be written on, so
    # "D) LAIN-LAIN: RONDA DEPARTMENT" is still that line. Only those lines get
    # this leeway; every other perkara must match by its whole name.
    fill_in = {norm for seq in standard for norm, label in zip(seq, labels_of[seq]['labels'])
               if label.rstrip().endswith(':')}
    for l in layouts:
        l['seq'] = tuple(next((f for f in fill_in if s.startswith(f)), s) for s in l['seq'])
    for l in layouts:
        if not l['marks'] or l['seq'] in standard or not standard:
            continue
        ratio = lambda ref: difflib.SequenceMatcher(a=ref, b=l['seq'], autojunk=False).ratio()
        ref = max(standard, key=ratio)
        form = labels_of[ref]['labels']
        if len(ref) == len(l['seq']):
            differ = [f"{i}: {got!r} (form: {want!r})" for i, (got, want, a, b)
                      in enumerate(zip(l['labels'], form, l['seq'], ref), 1) if a != b]
            warnings.append(f"{l['here']}: perkara names differ from the form at {'; '.join(differ)} - "
                            f"same number of rows, so scores kept in place")
            continue
        mapping, dropped, ok = [], [], True
        for op, a1, a2, b1, b2 in difflib.SequenceMatcher(a=ref, b=l['seq'], autojunk=False).get_opcodes():
            if op == 'equal' or (op == 'replace' and a2 - a1 == b2 - b1):
                mapping += range(b1, b2)
            elif op == 'delete':
                mapping += [None] * (a2 - a1)
            elif op == 'insert':
                dropped += range(b1, b2)
            else:
                ok = False
        if not ok:
            problems.add(l['here'], f"{len(l['seq'])} perkara rows where the form has {len(ref)}, "
                                    f"and they do not line up with it by name; fix the sheet")
            continue
        # A row the form lacks, with scores in it, is the same as a score typed
        # on a heading: the workbook's SUM counts it, but it is no perkara.
        scored = [j for j in dropped if any(m['scores'][j] is not None for m in l['marks'])]
        for m in l['marks']:
            m['scores'] = [None if j is None else m['scores'][j] for j in mapping]
        missing = [form[i] for i, j in enumerate(mapping) if j is None]
        warnings.append(f"{l['here']}: lined up with the form"
                        + (f"; dropped row(s) not on the form {[l['labels'][j] or '(no name)' for j in dropped]}"
                           if dropped else '')
                        + (" - they held scores, not counted (the workbook's JUMLAH does count them)" if scored else '')
                        + (f"; {missing} not on the sheet, counted as not marked" if missing else ''))
    return [labels_of[seq] for seq in standard]


SCORER_HEADERS = ['workbook', 'scorer on sheet', 'other spellings', 'weeks', 'staff marked (e.g.)',
                  'looks like a misspelling of', 'payroll_id']


def scorer_key(workbook, name):
    """Per workbook: the same name at two outlets may be two people."""
    return (workbook, ' '.join(name.split()).lower())


def write_scorer_list(path, marks):
    """Every scorer name on the sheets, one row per workbook and name (case
    merged), with an empty payroll_id column to fill in and read back with
    --scorer-ids."""
    groups = {}
    for m in marks:
        if m['scored_by']:
            g = groups.setdefault(scorer_key(m['workbook'], m['scored_by']),
                                  dict(spellings=Counter(), weeks=0, staff=[]))
            g['spellings'][' '.join(m['scored_by'].split())] += 1
            g['weeks'] += 1
            if m['person'] not in g['staff']:
                g['staff'].append(m['person'])
    wbx = openpyxl.Workbook()
    ws = wbx.active
    ws.title = 'scorers'
    ws.append(SCORER_HEADERS)
    for (book, low), g in sorted(groups.items()):
        # A much rarer spelling close to a common one at the same outlet is
        # probably the same person mistyped.
        likely = [o for (b, o), og in groups.items() if b == book and o != low and og['weeks'] > 3 * g['weeks']
                  and difflib.SequenceMatcher(a=low, b=o).ratio() >= 0.8]
        names = [n for n, _ in g['spellings'].most_common()]
        ws.append([book, names[0], ', '.join(names[1:]), g['weeks'], ', '.join(g['staff'][:4]),
                   ' / '.join(groups[(book, o)]['spellings'].most_common(1)[0][0] for o in likely), ''])
    for col, width in zip('ABCDEFG', (22, 26, 22, 8, 34, 28, 14)):
        ws.column_dimensions[col].width = width
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = 'A2'
    wbx.save(path)
    return len(groups)


def read_scorer_ids(path, problems):
    """(workbook, name) -> payroll number, from a filled-in --scorer-list file."""
    ws = openpyxl.load_workbook(path, data_only=True).active
    head = [text(c.value).lower() for c in ws[1]]
    try:
        book_i, name_i, id_i = head.index('workbook'), head.index('scorer on sheet'), head.index('payroll_id')
    except ValueError:
        sys.exit(f'{path}: the first row must keep the headings workbook, scorer on sheet and payroll_id')
    ids = {}
    for row in ws.iter_rows(min_row=2):
        book, name, pid = (text(row[i].value) for i in (book_i, name_i, id_i))
        if not (book and name and pid):
            continue
        pid = pid.replace(' ', '').upper()
        if pid == '-':  # known, and deliberately no scorer: saved without one
            ids[scorer_key(book, name)] = pid
            continue
        if not PAYROLL_RE.match(pid):
            problems.add(f'{path.name}!{row[id_i].coordinate}', f'{pid!r} is not a payroll number (letters then digits)')
            continue
        ids[scorer_key(book, name)] = pid
    return ids


def render(marks, standard, sources):
    out = []
    out.append('INSERT INTO import_marks (user_id, period_year, period_month, week_no, scored_by, note, scores) VALUES')
    rows, last_source = [], None
    for m in marks:
        if m['source'] != last_source:
            rows.append(f"  -- {comment(m['source'])}")
            last_source = m['source']
        for e in m['extra']:
            rows.append(f"  -- {m['person']} week {m['week']}: {comment(e)}")
        scores = ', '.join('NULL' if s is None else str(s) for s in m['scores'])
        rows.append(f"  ({sql_text(m['person'])}, {m['year']}, {m['month']}, {m['week']}, "
                    f"{sql_text(m['scored_by'])}, {sql_text(m['note'])}, ARRAY[{scores}])")
    # Commas go on the value rows only, never on the comment lines between them.
    value_idx = [i for i, r in enumerate(rows) if not r.lstrip().startswith('--')]
    for i in value_idx[:-1]:
        rows[i] += ','
    out.extend(rows)
    out.append(';')

    head = [
        f'-- GENERATED {dt.datetime.now():%Y-%m-%d %H:%M} by xlsx_to_import_marks.py from:',
        *[f'--   {s}' for s in sources],
        f'-- {len(marks)} week-marks for {len({m["person"] for m in marks})} people.',
        '-- Perkara order of the form (score list position -> row label):',
    ]
    for layout in standard:
        head.append(f"--   {len(layout['labels'])}-line form, as on {comment(layout['here'])}:")
        head.extend(f'--     {i:2d} {comment(l)}' for i, l in enumerate(layout['labels'], 1))
    return head, '\n'.join(out)


def main():
    for stream in (sys.stdout, sys.stderr):  # a sheet name may not fit the console's code page
        stream.reconfigure(errors='replace')
    # openpyxl complains about workbook features it does not read (data
    # validation, a date cell holding garbage). Neither touches a score; a bad
    # cell that matters is reported below with its address.
    pywarnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')
    ap = argparse.ArgumentParser(description=__doc__.split('\n\n')[0])
    ap.add_argument('inputs', nargs='+', type=Path, help='.xlsx files or folders of them')
    ap.add_argument('--month', help='only this month, e.g. 2026-09')
    ap.add_argument('--until', help='only up to and including this month, e.g. 2026-09')
    ap.add_argument('-o', '--out', type=Path, help='where to write the SQL (default: import_marks_ready.sql beside the first input)')
    ap.add_argument('--split-months', action='store_true',
                    help='one file per month (OUT_2026-01.sql, ...); the Supabase SQL editor refuses a file much over 1 MB')
    ap.add_argument('--skip-weeks-of', nargs='+', metavar='SQL',
                    help='earlier import files (wildcards allowed): leave out any person-week they already hold, '
                         'since a person has one mark a week')
    ap.add_argument('--scorer-list', type=Path,
                    help='also write every scorer name to this .xlsx, with a payroll_id column to fill in')
    ap.add_argument('--scorer-ids', type=Path,
                    help='a filled-in --scorer-list file: each name with a payroll_id is replaced by it; "-" means no scorer')
    ap.add_argument('--check-file', action='store_true',
                    help='also write OUT_check.sql: one run that lists every problem across all the rows and saves nothing')
    args = ap.parse_args()

    def month_arg(flag, value):
        m = re.fullmatch(r'(\d{4})-(\d{1,2})', value or '')
        if value and not m:
            sys.exit(f'{flag} must look like 2026-09')
        return (int(m[1]), int(m[2])) if m else None

    only_month = month_arg('--month', args.month)
    until = month_arg('--until', args.until)

    files = []
    for p in args.inputs:
        if p.is_dir():
            files += sorted(x for x in p.glob('*.xlsx') if not x.name.startswith('~$'))
        elif p.exists():
            files.append(p)
        else:
            sys.exit(f'not found: {p}')
    if not files:
        sys.exit('no .xlsx files found')

    problems, warnings, marks, layouts = Problems(), [], [], []
    for f in files:
        read_workbook(f, only_month, until, problems, warnings, marks, layouts)
    standard = line_up(layouts, problems, warnings)

    if args.skip_weeks_of:
        # A person has one mark a week; weeks an earlier import already holds
        # (staff marks for someone also on the SV/AS sheets) stay as they are.
        taken = set()
        for pattern in args.skip_weeks_of:
            paths = sorted(glob.glob(pattern)) or sys.exit(f'--skip-weeks-of: nothing matches {pattern}')
            for p in paths:
                for line in open(p, encoding='utf-8'):
                    m = re.match(r"\s*\('([^']*)', (\d+), (\d+), (\d), ", line)
                    if m:
                        taken.add((m[1].upper(), int(m[2]), int(m[3]), int(m[4])))
        kept = [m for m in marks if (m['person'], m['year'], m['month'], m['week']) not in taken]
        left_out = Counter(m['person'] for m in marks if (m['person'], m['year'], m['month'], m['week']) in taken)
        print(f'{len(marks) - len(kept)} weeks left out - already in the earlier import files: '
              + ', '.join(f'{p} ({n})' for p, n in sorted(left_out.items())))
        marks[:] = kept

    if args.scorer_list:
        print(f'{write_scorer_list(args.scorer_list, marks)} scorer names written to {args.scorer_list}')
    if args.scorer_ids:
        ids = read_scorer_ids(args.scorer_ids, problems)
        swapped, still = 0, Counter()
        for m in marks:
            if m['scored_by']:
                pid = ids.get(scorer_key(m['workbook'], m['scored_by']))
                if pid:
                    m['scored_by'] = None if pid == '-' else pid
                    swapped += 1
                else:
                    still[(m['workbook'], m['scored_by'])] += 1
        print(f'{swapped} weeks take their scorer from {args.scorer_ids.name} (payroll number, or "-" for none); '
              f'{sum(still.values())} weeks ({len(still)} names) have no row there, so their name is matched at the outlet:')
        for (book, name), n in sorted(still.items()):
            print(f'  {book}: {name!r} ({n} weeks)')

    for w in warnings:
        print('note:', w)
    if problems:
        print(f'\nNothing written. Fix these {len(problems)} cells and run again:', file=sys.stderr)
        for p in problems:
            print('  ' + p, file=sys.stderr)
        sys.exit(1)
    if not marks:
        sys.exit('no scored weeks found' + (f' for {args.month}' if args.month else ''))

    template = TEMPLATE.read_text(encoding='utf-8')
    block = re.compile(r'^INSERT INTO import_marks \(.*?^;[ \t]*$', re.M | re.S)
    if not block.search(template):
        sys.exit(f'could not find the INSERT INTO import_marks ... ; block in {TEMPLATE}')
    out = args.out or files[0].with_name('import_marks_ready.sql')
    parts = {out: marks}
    if args.split_months:
        parts = {}
        for m in marks:
            parts.setdefault(out.with_name(f"{out.stem}_{m['year']}-{m['month']:02d}{out.suffix}"), []).append(m)

    per_month = Counter((m['year'], m['month']) for m in marks)
    print(f'\n{len(marks)} week-marks for {len({m["person"] for m in marks})} people:')
    for (y, mo), n in sorted(per_month.items()):
        print(f'  {y}-{mo:02d}: {n}')

    print('\nWritten:')
    for path, part in sorted(parts.items()):
        head, values = render(part, standard, [str(f) for f in files])
        sql = '\n'.join(head) + '\n\n' + block.sub(lambda _: values, template, count=1)
        path.write_text(sql, encoding='utf-8')
        size = path.stat().st_size
        print(f'  {path}  ({size // 1024} KB, {len(part)} week-marks)')
        if size > 900_000:
            print('    note: the Supabase SQL editor refuses files much over 1 MB - run again with --split-months')

    if args.check_file:
        # The same SQL with check_only on, so it judges by exactly the import's
        # rules. Problems hang on the person and the scorer, not the week, so
        # one week per person, scorer and layout is enough to find them all —
        # plus any week with a score out of range, which only that week shows.
        seen, sample = set(), []
        for m in marks:
            key = (m['person'], (m['scored_by'] or '').strip().lower(), len(m['scores']))
            if key not in seen or any(s is not None and not 0 <= s <= 5 for s in m['scores']):
                seen.add(key)
                sample.append(m)
        switch = 'false        AS check_only;'
        if switch not in template:
            sys.exit(f'could not find the check_only setting in {TEMPLATE}')
        head, values = render(sample, standard, [str(f) for f in files])
        head.insert(0, f'-- CHECK FILE: one week per person and scorer ({len(sample)} of {len(marks)}). '
                       f'Lists every problem; saves nothing.')
        sql = '\n'.join(head) + '\n\n' + block.sub(lambda _: values, template.replace(switch, 'true         AS check_only;'), count=1)
        path = out.with_name(f'{out.stem}_check{out.suffix}')
        path.write_text(sql, encoding='utf-8')
        print(f'  {path}  ({path.stat().st_size // 1024} KB, check only - run this first)')
    print('Paste each into the Supabase SQL editor and run it.')


if __name__ == '__main__':
    main()
